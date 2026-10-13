# hand check workbook
import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "hand_check.xlsx"
if OUT.exists():
    sys.exit(f"{OUT.name} exists; move it away first.")

HEAD = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="44546A")
INPUT = PatternFill("solid", fgColor="FFF2CC")
BODY = Font(name="Arial")
GREEK = Font(name="Arial", size=12)
RTL = Alignment(horizontal="right", readingOrder=2, wrap_text=True, vertical="top")
WRAP = Alignment(wrap_text=True, vertical="top")


def read(path):
    with open(ROOT / path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def sheet(wb, title, cols, rows, inputs, rtl=(), greek=("Greek",), choices="y,n"):
    # cols: (name, width)
    ws = wb.create_sheet(title)
    names = [c for c, _ in cols] + inputs
    for i, name in enumerate(names, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
        width = dict(cols).get(name, 14)
        ws.column_dimensions[c.column_letter].width = width
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = GREEK if names[i - 1] in greek else BODY
            c.alignment = RTL if names[i - 1] in rtl else WRAP
        for i in range(len(row) + 1, len(names) + 1):
            ws.cell(row=r, column=i).fill = INPUT
    last = len(rows) + 1
    if choices and rows:
        dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
        ws.add_data_validation(dv)
        for i, name in enumerate(names, 1):
            if name in inputs and name != "Notes":
                col = ws.cell(row=1, column=i).column_letter
                dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
start = wb.active
start.title = "Start here"
start.column_dimensions["A"].width = 110
for i, line in enumerate([
    "Dragoman: hand check",
    "",
    "Fill in the yellow columns only (y / n, or a note). Work in this order:",
    "1 Traps: is each trap sound? LSJ ok = Greek lacks the 'via English' sense. Persian ok = both Persian columns right.",
    "2 P2 hits: every answer the script flagged as a trap firing. Real P2 = it really is the English-only sense.",
    "3 Findings: a hand-read of the direct run. Agree = you see the same thing in the raw answers.",
    "4 Temperature: distinct answers flagged for the temperature items. Real = the effect is really there.",
    "5 Kinship key: are the Persian side-kept / side-lost words right?",
    "6 Kinship answers: distinct answers with the automatic class. Your class = kept / both / wrong / generic / other.",
    "   Sorted by class: kept, both, wrong, generic first (the counted ones); 'other' is a quick skim for misses.",
    "",
    "Sheets 4 and 6 list each distinct answer once; 'Times' says how often it occurred.",
    "Automatic flags come from score.py; full rows are in results/*.csv.",
]):
    start.cell(row=i + 1, column=1, value=line).font = Font(name="Arial", bold=(i == 0), size=14 if i == 0 else 11)

key = {r["id"]: r for r in read("data/key.csv")}

# 1 traps
traps = [r for r in key.values() if r["role"] == "trap"]
sheet(wb, "1 Traps",
      [("ID", 6), ("Greek", 14), ("POS", 9), ("Sense", 22), ("English", 12), ("Persian correct", 20),
       ("Persian via English", 20), ("Paired with", 12), ("Note", 36)],
      [[int(r["id"]), r["greek"], r["pos"], r["sense"], r["english"], r["fa_correct"], r["fa_via_english"],
        r["paired_with"], r["note"]] for r in traps],
      ["LSJ ok", "Persian ok", "Keep", "Notes"], rtl=("Persian correct", "Persian via English"))

# 2 P2 hits
hits = [r for r in read("results/words.csv") if r["p2_found"] == "yes"]
sheet(wb, "2 P2 hits",
      [("ID", 6), ("Greek", 14), ("Sense", 18), ("English", 10), ("Model", 14), ("Condition", 10), ("Sample", 7),
       ("English step 1", 12), ("Answer", 36), ("Persian via English", 18)],
      [[int(r["id"]), r["greek"], key[r["id"]]["sense"], key[r["id"]]["english"], r["model"], r["condition"],
        int(r["sample"]), r["english_step1"], r["equivalents"], key[r["id"]]["fa_via_english"]] for r in hits],
      ["Real P2", "Notes"], rtl=("Answer", "Persian via English"))

# 3 findings
sheet(wb, "3 Findings",
      [("Category", 22), ("ID", 6), ("Greek", 14), ("Models", 18), ("What happened", 40), ("Samples", 14),
       ("English step 1", 22), ("Note", 36)],
      [[r["category"], int(r["id"]), r["greek"], r["models"], r["what_happened"], r["samples"],
        r["english_pivot_step1"], r["note"]] for r in read("results/findings.csv")],
      ["Agree", "Notes"])


def distinct(rows, keyf):
    # one row per distinct answer
    groups = defaultdict(list)
    for r in rows:
        groups[keyf(r)].append(r)
    return groups


# 4 temperature
temp = distinct([r for r in read("results/temperature.csv") if r["effect_seen"] == "yes"],
                lambda r: (int(r["id"]), r["greek"], r["effect"], r["equivalents"]))
sheet(wb, "4 Temperature",
      [("ID", 6), ("Greek", 12), ("Effect", 16), ("Answer", 40), ("Times", 7), ("Models", 22), ("Temperatures", 16)],
      [[*k, len(v), ", ".join(sorted({r["model"] for r in v})), ", ".join(sorted({r["temperature"] for r in v}))]
       for k, v in sorted(temp.items())],
      ["Real", "Notes"], rtl=("Answer",))

# 5 kinship key
sheet(wb, "5 Kinship key",
      [("ID", 6), ("Greek", 14), ("LSJ gloss", 30), ("English", 14), ("Persian side kept", 22),
       ("Persian side lost", 24), ("Strength", 9), ("Note", 36)],
      [[int(r["id"]), r["greek"], r["lsj_gloss"], r["english"], r["fa_side_kept"], r["fa_side_lost"],
        r["strength"], r["note"]] for r in read("data/kinship_key.csv")],
      ["Persian ok", "Notes"], rtl=("Persian side kept", "Persian side lost"))

# 6 kinship answers
order = {c: i for i, c in enumerate(["kept", "both", "wrong", "generic", "other", "missing"])}
kin = distinct(read("results/kinship.csv"), lambda r: (int(r["id"]), r["greek"], r["equivalents"], r["class"]))
sheet(wb, "6 Kinship answers",
      [("ID", 6), ("Greek", 14), ("Answer", 40), ("Auto class", 10), ("Times", 7), ("Models", 24), ("Conditions", 14)],
      [[*k, len(v), ", ".join(sorted({r["model"] for r in v})), ", ".join(sorted({r["condition"] for r in v}))]
       for k, v in sorted(kin.items(), key=lambda kv: (order[kv[0][3]], kv[0][0], kv[0][2]))],
      ["Your class", "Notes"], rtl=("Answer",), choices="kept,both,wrong,generic,other")

wb.save(OUT)
print(OUT.name, {ws.title: ws.max_row - 1 for ws in wb.worksheets[1:]})
